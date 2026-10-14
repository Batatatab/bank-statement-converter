from datetime import date, datetime
import re

MONTH_NAME_TO_NUM = {
    "jan": 1,
    "january": 1,
    "feb": 2,
    "february": 2,
    "mar": 3,
    "march": 3,
    "apr": 4,
    "april": 4,
    "may": 5,
    "jun": 6,
    "june": 6,
    "jul": 7,
    "july": 7,
    "aug": 8,
    "august": 8,
    "sep": 9,
    "sept": 9,
    "september": 9,
    "oct": 10,
    "october": 10,
    "nov": 11,
    "november": 11,
    "dec": 12,
    "december": 12,
}

DECIMAL_REGEX = re.compile(r"^[-+]?[\d,]+\.\d{2}$")
CHEQUE_6DIGIT_REGEX = re.compile(r"^\d{6}$")

FOOTER_AND_TOTAL_PATTERNS = [
    r"\bPERIOD\s*TOTAL\b",
    r"\bGRAND\s*TOTAL\b",
    r"\bSUB\s*TOTAL\b",
    r"\bPAGE\s*TOTAL\b",
    r"\bTOTAL\s*DEBIT\b",
    r"\bTOTAL\s*CREDIT\b",
    r"\bTOTAL\s*WITHDRAWAL\b",
    r"\bTOTAL\s*DEPOSIT\b",
    r"\bSTATEMENT\s*PERIOD\b",
    r"\bSTATEMENT\s*OF\s*ACCOUNT\b",
    r"\bCUSTOMER\s*ID\b",
    r"\bACCOUNT\s*NO\b",
    r"\bREGISTERED\s*OFFICE\b",
    r"\bPAGE\s*(?:NO)?\.?\s*\d+\s*(?:OF\s*\d+)?",
    r"\bCOMPUTER\s*GENERATED\b",
    r"\bEND\s*OF\s*STATEMENT\b",
    r"\bHTML\s*TO\s*PDF\b",
    r"\bDEVELOPER-FRIENDLY\b",
    r"\bEXPLORE\s*OUR\b",
]

ANNEXURE_KEYWORDS = [
    r"\bACCOUNT\s*RELATED\s*OTHER\s*INFORMATION\b",
    r"\bLEGENDS\s*(?:FOR\s*TRANSACTIONS)?\b",
    r"\bNOMINEE\s*NAME\b",
    r"\bAUTHORISED\s*SIGNATORY\b",
    r"\bINCOME\s*TAX\s*ACT\b",
    r"\bFORM\s*15[GH]\b",
    r"\bVAT/MAT/NFS\b",
    r"\bCASH\s*WITHDRAWAL\s*AT\s*OTHER\s*BANK\b",
    r"\bINTERNET\s*FUNDS\s*TRANSFER\b",
    r"\bREAL\s*TIME\s*GROSS\s*SETTLEMENT\b",
    r"\bNATIONAL\s*ELECTRONIC\s*FUND\s*TRANSFER\b",
    r"\bTERMS\s*(?:AND|&)\s*CONDITIONS\b",
    r"\bIMPORTANT\s*INFORMATION\b",
    r"\bCUSTOMER\s*CARE\b",
    r"\bPLEASE\s*NOTE\b",
    r"\bPLEASE\s*DO\s*NOT\s*SHARE\b",
    r"\bUNLESS\s*A\s*CONSTITUENT\b",
]

OPENING_BAL_PATTERNS = [
    r"\bop(?:en(?:ing)?)?\.?\s*bal(?:ance)?\.?\b",
    r"\b(?:to\s+|by\s+)?bal(?:ance)?\.?\s*b[/\.]?d\b",
    r"\b(?:to\s+|by\s+)?bal(?:ance)?\.?\s*b[/\.]?f(?:wd)?\b",
    r"\bb[/\.]?[df](?:wd)?\s*bal(?:ance)?\.?\b",
    r"\b(?:to\s+|by\s+)?(?:bal(?:ance)?\.?\s*)?brought\s*(?:down|forward)\b",
    r"\bbalance\s*forward(?:ed)?\b",
    r"\b(?:start(?:ing)?|initial|opening\s*book)\s*bal(?:ance)?\.?\b",
]

CLOSING_BAL_PATTERNS = [
    r"\bcl(?:o(?:s(?:e|ing)?)?)?\.?\s*bal(?:ance)?\.?\b",
    r"\b(?:to\s+|by\s+)?bal(?:ance)?\.?\s*c[/\.]?d\b",
    r"\b(?:to\s+|by\s+)?bal(?:ance)?\.?\s*c[/\.]?f(?:wd)?\b",
    r"\bc[/\.]?[df](?:wd)?\s*bal(?:ance)?\.?\b",
    r"\b(?:to\s+|by\s+)?(?:bal(?:ance)?\.?\s*)?carried\s*(?:down|forward)\b",
    r"\b(?:end(?:ing)?|final|closing\s*book)\s*bal(?:ance)?\.?\b",
]

HEADER_KEYWORDS = {
    "serial_no": ["sl no", "sr no", "s.no", "sr.no", "sl.no", "s no"],
    "val_date": [
        "value date",
        "value dt",
        "val date",
        "val dt",
        "value_date",
        "v.date",
    ],
    "date": [
        "transaction date",
        "txn date",
        "tran date",
        "trans date",
        "date",
        "post date",
    ],
    "chq_ref": [
        "reference no",
        "ref no",
        "ref. no.",
        "reference",
        "cheque/reference#",
        "cheque number",
        "chq no",
        "cheque no",
        "cheque",
        "chq / ref",
    ],
    "particulars": [
        "description",
        "particulars",
        "narration",
        "transaction details",
        "remarks",
        "details",
    ],
    "dr": [
        "debit (receipts)",
        "debited amount",
        "debited",
        "debit",
        "withdrawal",
        "dr amount",
        "withdrawals",
        "receipts",
        "dr",
    ],
    "cr": [
        "credit (payments)",
        "credited amount",
        "credited",
        "credit",
        "deposit",
        "cr amount",
        "deposits",
        "payments",
        "cr",
    ],
    "bal": ["balance", "bal", "closing balance", "running balance"],
}


def _match_keyword(text: str, keywords: list[str]) -> bool:
  if not text:
    return False
  t_clean = text.lower().strip()
  for kw in keywords:
    pattern = r"\b" + re.escape(kw.lower()) + r"\b"
    if re.search(pattern, t_clean):
      return True
  return False


def is_opening_balance_line(text: str) -> bool:
  if not text:
    return False
  return any(
      re.search(pat, text, re.IGNORECASE) for pat in OPENING_BAL_PATTERNS
  )


def is_closing_balance_line(text: str) -> bool:
  if not text:
    return False
  return any(
      re.search(pat, text, re.IGNORECASE) for pat in CLOSING_BAL_PATTERNS
  )


def is_annexure_or_legend_line(text: str) -> bool:
  if not text:
    return False
  return any(re.search(pat, text, re.IGNORECASE) for pat in ANNEXURE_KEYWORDS)


def is_total_summary_line(
    line_full_text: str, line_words: list, decimal_items: list
) -> bool:
  if not line_full_text:
    return False
  clean = line_full_text.strip().upper()

  if re.match(
      r"^(?:GRAND\s+|PERIOD\s+|SUB\s+|PAGE\s+|STATEMENT\s+)?TOTAL(?:\s*[:\-])?\b",
      clean,
  ):
    return True

  first_word = line_words[0]["text"].strip().upper() if line_words else ""
  if first_word in ("TOTAL", "TOTAL:", "TOTAL-"):
    return True

  if re.search(r"\bTOTAL\b", clean) and len(decimal_items) >= 2:
    if not re.search(
        r"\b(?:PVT|LTD|STORE|PETROL|PUMP|SERVICES|EXPENSE|PAYMENT)\b", clean
    ):
      return True

  return False


def clean_float(val) -> float:
  if val is None:
    return 0.0
  s = str(val).strip()
  if not s or s == "-":
    return 0.0
  match = re.search(r"[-+]?[\d,]+\.\d{2,3}(?:\s*(?:CR|DR))?", s, re.IGNORECASE)
  if not match:
    return 0.0
  target = match.group(0).replace(",", "").strip()
  is_negative = (
      target.startswith("-") or "(DR)" in target.upper() or "OD" in target.upper()
  )
  s_clean = re.sub(r"[^\d\.]", "", target)
  try:
    f = float(s_clean)
    return -f if is_negative else f
  except ValueError:
    return 0.0


def extract_strict_decimal(text: str) -> float:
  c = (
      str(text)
      .replace(",", "")
      .replace("+", "")
      .replace("₹", "")
      .replace("Rs.", "")
      .strip()
  )
  c = re.sub(r"(?:CR|DR)$", "", c, flags=re.IGNORECASE).strip()
  if re.match(r"^[-+]?\d+\.\d{2,3}$", c):
    try:
      return abs(round(float(c), 2))
    except ValueError:
      return 0.0
  return 0.0


def parse_date_flexible(val, default_year: int = None) -> date | None:
  if not val:
    return None
  if isinstance(val, (datetime, date)):
    return val.date() if isinstance(val, datetime) else val
  s = str(val).strip()
  if not s:
    return None
  s = re.sub(r"[T\s]\d{1,2}:\d{2}(?::\d{2})?.*$", "", s).strip().rstrip(",")

  def make_date(y, m, d):
    if y < 70:
      y += 2000
    elif y < 100:
      y += 1900
    try:
      return date(y, m, d)
    except (ValueError, TypeError):
      return None

  m1 = re.search(r"\b(\d{1,2})[/\.\- ]([A-Za-z]+)[/\.\- ](\d{2,4})\b", s)
  if m1:
    d_v = int(m1.group(1))
    m_v = m1.group(2).lower()
    y_v = int(m1.group(3))
    if m_v in MONTH_NAME_TO_NUM:
      res = make_date(y_v, MONTH_NAME_TO_NUM[m_v], d_v)
      if res:
        return res

  m2 = re.search(
      r"\b([A-Za-z]+)[/\.\- ](\d{1,2})(?:st|nd|rd|th)?,?[/\.\- ](\d{2,4})\b", s
  )
  if m2:
    m_v = m2.group(1).lower()
    d_v = int(m2.group(2))
    y_v = int(m2.group(3))
    if m_v in MONTH_NAME_TO_NUM:
      res = make_date(y_v, MONTH_NAME_TO_NUM[m_v], d_v)
      if res:
        return res

  m3 = re.search(r"\b(\d{4})[/\.\-](\d{1,2})[/\.\-](\d{1,2})\b", s)
  if m3:
    res = make_date(int(m3.group(1)), int(m3.group(2)), int(m3.group(3)))
    if res:
      return res

  m4 = re.search(r"\b(\d{1,2})[/\.\-](\d{1,2})[/\.\-](\d{2,4})\b", s)
  if m4:
    p1, p2, p3 = int(m4.group(1)), int(m4.group(2)), int(m4.group(3))
    res = make_date(p3, p2, p1)
    if res:
      return res

  m5 = re.search(r"\b(\d{1,2})\s+(\d{1,2})\s+(\d{2,4})\b", s)
  if m5:
    p1, p2, p3 = int(m5.group(1)), int(m5.group(2)), int(m5.group(3))
    if 1 <= p2 <= 12 and 1 <= p1 <= 31:
      res = make_date(p3, p2, p1)
      if res:
        return res

  if default_year:
    m6 = re.search(r"\b(\d{1,2})[/\.\- ]([A-Za-z]+)\b", s)
    if m6 and m6.group(2).lower() in MONTH_NAME_TO_NUM:
      res = make_date(
          default_year, MONTH_NAME_TO_NUM[m6.group(2).lower()], int(m6.group(1))
      )
      if res:
        return res

    m7 = re.search(r"\b([A-Za-z]+)[/\.\- ](\d{1,2})\b", s)
    if m7 and m7.group(1).lower() in MONTH_NAME_TO_NUM:
      res = make_date(
          default_year, MONTH_NAME_TO_NUM[m7.group(1).lower()], int(m7.group(2))
      )
      if res:
        return res

    m8 = re.search(r"\b(\d{1,2})[/\-](\d{1,2})\b", s)
    if m8:
      p1, p2 = int(m8.group(1)), int(m8.group(2))
      if 1 <= p2 <= 12 and 1 <= p1 <= 31:
        res = make_date(default_year, p2, p1)
        if res:
          return res

  return None


class AnchorStatementParser:

  def __init__(self, log_fn=None):
    self.log = log_fn or (lambda x: None)

  def _detect_dates_in_line(self, line, default_year=None):
    found = []
    consumed = set()
    n = min(14, len(line))

    idx = 0
    while idx < n:
      if idx in consumed:
        idx += 1
        continue

      matched = False
      for k in [3, 2, 1]:
        if idx + k <= len(line):
          span_words = line[idx : idx + k]
          span_text = " ".join(w["text"].strip() for w in span_words)
          d = parse_date_flexible(span_text, default_year=default_year)
          if d:
            span_indices = set(range(idx, idx + k))
            found.append((d, span_text, span_indices))
            consumed.update(span_indices)
            idx += k
            matched = True
            break
      if not matched:
        idx += 1

    return found

  def parse(
      self,
      pdf_path: str,
      password: str = None,
      first_tx_date_str: str = "",
      opening_bal: float = None,
      expected_closing_bal: float = None,
  ):
    import pandas as pd
    import pdfplumber

    target_date = parse_date_flexible(first_tx_date_str)
    if not target_date:
      raise ValueError(
          f"Could not parse First Transaction Date: '{first_tx_date_str}'"
      )

    with pdfplumber.open(
        pdf_path, password=password if password else None
    ) as pdf:
      self.log(
          f"Opened PDF: {len(pdf.pages)} pages. Locating anchor:"
          f" {target_date.strftime('%d-%m-%Y')}..."
      )

      anchor_found, anchor_page_idx, header_cols, anchor_y = (
          self._locate_anchor_and_headers(pdf, target_date)
      )
      if not anchor_found:
        raise ValueError(
            f"Could not locate transaction for date '{first_tx_date_str}'."
            " Verify statement input."
        )

      self.log(
          f"Identified column boundaries on Page {anchor_page_idx}:"
          f" {list(header_cols.keys())}"
      )
      raw_txs, in_table_op_bal, in_table_cl_bal = (
          self._extract_all_transactions_borderless(
              pdf, anchor_page_idx, anchor_y, header_cols, target_date.year
          )
      )

    self.log(f"Extracted {len(raw_txs)} legitimate transactions.")

    effective_op_bal = opening_bal
    if effective_op_bal is None and in_table_op_bal is not None:
      effective_op_bal = in_table_op_bal
      self.log(f"Adopted in-table Opening Balance: ₹{effective_op_bal:,.2f}")

    # Reverse chronological sequence handling
    parsed_dates = [tx["date_obj"] for tx in raw_txs if tx.get("date_obj")]
    if len(parsed_dates) >= 2:
      desc_count = sum(
          1
          for i in range(len(parsed_dates) - 1)
          if parsed_dates[i] > parsed_dates[i + 1]
      )
      asc_count = sum(
          1
          for i in range(len(parsed_dates) - 1)
          if parsed_dates[i] < parsed_dates[i + 1]
      )
      if desc_count > asc_count:
        self.log(
            "Detected reverse chronology (newest first). Inverting to"
            " chronological sequence..."
        )
        raw_txs.reverse()

    reconciled_rows, calculated_op_bal, running_math_bal, is_od_account = (
        self._reconcile(raw_txs, effective_op_bal)
    )

    last_extracted_pdf_bal = None
    for tx in reversed(raw_txs):
      if tx.get("bal") and tx["bal"] != 0.0:
        last_extracted_pdf_bal = tx["bal"]
        break

    if in_table_cl_bal is not None and in_table_cl_bal != 0.0:
      definitive_closing_bal = in_table_cl_bal
      self.log(
          "Definitive Statement Closing Balance (from Bal c/d / Summary):"
          f" ₹{definitive_closing_bal:,.2f}"
      )
    elif last_extracted_pdf_bal is not None:
      definitive_closing_bal = last_extracted_pdf_bal
      self.log(
          "Definitive Statement Closing Balance (from Last Tx PDF Bal):"
          f" ₹{definitive_closing_bal:,.2f}"
      )
    else:
      definitive_closing_bal = (
          reconciled_rows[-1]["Bal"] if reconciled_rows else calculated_op_bal
      )

    final_bal = definitive_closing_bal
    total_dr = sum(r["Dr"] for r in reconciled_rows[1:])
    total_cr = sum(r["Cr"] for r in reconciled_rows[1:])

    target_to_compare = (
        expected_closing_bal
        if expected_closing_bal is not None
        else definitive_closing_bal
    )
    diff = round(final_bal - target_to_compare, 2)
    closing_matches = abs(diff) <= 0.05

    math_diff = round(running_math_bal - final_bal, 2)
    math_reconciled = abs(math_diff) <= 0.05

    verification_info = {
        "opening_bal": calculated_op_bal,
        "final_bal": final_bal,
        "running_math_bal": running_math_bal,
        "total_dr": total_dr,
        "total_cr": total_cr,
        "tx_count": max(0, len(reconciled_rows) - 1),
        "expected_closing_bal": target_to_compare,
        "closing_matches": closing_matches,
        "math_reconciled": math_reconciled,
        "diff": diff,
        "math_diff": math_diff,
        "is_od_account": is_od_account,
    }

    return pd.DataFrame(reconciled_rows), verification_info

  def _locate_anchor_and_headers(self, pdf, target_date):
    for page_idx, page in enumerate(pdf.pages, start=1):
      words = page.extract_words(x_tolerance=2, y_tolerance=2)
      if not words:
        continue

      lines = self._group_words_by_line(words)

      for line_idx, line in enumerate(lines):
        dates_found = self._detect_dates_in_line(
            line, default_year=target_date.year
        )
        detected_dates = [df_item[0] for df_item in dates_found]
        has_decimals = any(extract_strict_decimal(w["text"]) > 0 for w in line)

        if target_date in detected_dates and has_decimals:
          anchor_y = line[0]["top"]

          header_words = []
          header_line_idx = None
          for prev_idx in range(line_idx - 1, max(-1, line_idx - 15), -1):
            prev_line = lines[prev_idx]
            prev_str = " ".join([w["text"].lower() for w in prev_line])
            if any(
                _match_keyword(prev_str, HEADER_KEYWORDS[k])
                for k in ["dr", "cr", "bal", "particulars"]
            ):
              header_words.extend(prev_line)
              header_line_idx = prev_idx

          if header_line_idx is not None:
            for mid_idx in range(header_line_idx + 1, line_idx):
              mid_text = " ".join(w["text"] for w in lines[mid_idx])
              if is_opening_balance_line(mid_text):
                anchor_y = lines[mid_idx][0]["top"]
                break

          header_cols = self._build_header_intervals(header_words, page.width)
          return True, page_idx, header_cols, anchor_y

    return False, 0, {}, 0.0

  def _build_header_intervals(self, header_words, page_width):
    col_positions = {}
    for key in [
        "val_date",
        "date",
        "particulars",
        "chq_ref",
        "dr",
        "cr",
        "bal",
    ]:
      best_match = None
      for w in header_words:
        if _match_keyword(w["text"], HEADER_KEYWORDS[key]):
          best_match = w
          break
      if best_match:
        col_positions[key] = (best_match["x0"] + best_match["x1"]) / 2

    if "bal" not in col_positions:
      col_positions["bal"] = page_width * 0.88
    if "cr" not in col_positions:
      col_positions["cr"] = col_positions["bal"] - (page_width * 0.12)
    if "dr" not in col_positions:
      col_positions["dr"] = col_positions["cr"] - (page_width * 0.12)
    if "particulars" not in col_positions:
      col_positions["particulars"] = min(
          col_positions["dr"], col_positions["cr"]
      ) - (page_width * 0.25)
    if "date" not in col_positions:
      col_positions["date"] = page_width * 0.06

    sorted_cols = sorted(col_positions.items(), key=lambda item: item[1])
    intervals = {}
    for i in range(len(sorted_cols)):
      col_name = sorted_cols[i][0]
      curr_center = sorted_cols[i][1]
      left_bound = 0 if i == 0 else (sorted_cols[i - 1][1] + curr_center) / 2
      right_bound = (
          page_width
          if i == len(sorted_cols) - 1
          else (curr_center + sorted_cols[i + 1][1]) / 2
      )
      intervals[col_name] = (left_bound, right_bound)

    return intervals

  def _group_words_by_line(self, words):
    words_sorted = sorted(words, key=lambda w: (round(w["top"], 1), w["x0"]))
    lines = []
    curr = []
    last_top = None

    for w in words_sorted:
      if last_top is None or abs(w["top"] - last_top) <= 3.5:
        curr.append(w)
        last_top = w["top"]
      else:
        lines.append(sorted(curr, key=lambda item: item["x0"]))
        curr = [w]
        last_top = w["top"]
    if curr:
      lines.append(sorted(curr, key=lambda item: item["x0"]))
    return lines

  def _extract_all_transactions_borderless(
      self, pdf, start_page_idx, start_y, header_cols, default_year
  ):
    raw_txs = []
    current_tx = None
    in_table_op_bal = None
    in_table_cl_bal = None
    statement_ended = False

    dr_left, dr_right = header_cols.get("dr", (0, 0))
    cr_left, cr_right = header_cols.get("cr", (0, 0))
    bal_left, bal_right = header_cols.get("bal", (0, 0))
    part_left, part_right = header_cols.get("particulars", (0, 0))

    dr_center = (dr_left + dr_right) / 2
    cr_center = (cr_left + cr_right) / 2
    bal_center = (bal_left + bal_right) / 2
    amount_col_left_boundary = min(dr_left, cr_left)

    for page_idx in range(start_page_idx - 1, len(pdf.pages)):
      if statement_ended:
        break

      page = pdf.pages[page_idx]
      words = page.extract_words(x_tolerance=2, y_tolerance=2)
      if not words:
        continue

      lines = self._group_words_by_line(words)

      for line in lines:
        if page_idx == start_page_idx - 1 and line[0]["top"] < start_y - 3:
          continue

        line_full_text = " ".join([w["text"] for w in line]).strip()
        if not line_full_text:
          continue

        if is_annexure_or_legend_line(line_full_text):
          if len(raw_txs) > 0 or current_tx is not None:
            statement_ended = True
            break
          continue

        dates_found = self._detect_dates_in_line(
            line, default_year=default_year
        )
        date_indices = set()
        line_date = None
        date_str = ""

        if dates_found:
          line_date = dates_found[0][0]
          date_str = dates_found[0][1]
          for d_item in dates_found:
            date_indices.update(d_item[2])

        decimal_items = []
        particulars_tokens = []

        for w_idx, w in enumerate(line):
          val = extract_strict_decimal(w["text"])
          mid_x = (w["x0"] + w["x1"]) / 2
          if val > 0:
            decimal_items.append((val, mid_x))
          else:
            if w_idx in date_indices:
              continue
            # Universal particulars assignment (supports particulars on left or far right)
            if part_left > 0 or part_right > 0:
              if part_left <= mid_x <= part_right:
                particulars_tokens.append(w["text"])
            else:
              if mid_x < amount_col_left_boundary or mid_x > bal_right:
                particulars_tokens.append(w["text"])

        if is_total_summary_line(line_full_text, line, decimal_items):
          if len(decimal_items) >= 1:
            cl_cand = decimal_items[-1][0]
            if cl_cand > 0 and in_table_cl_bal is None:
              in_table_cl_bal = cl_cand
              self.log(
                  "Captured Closing Balance from Statement TOTAL:"
                  f" ₹{cl_cand:,.2f}"
              )
          clean_upper = line_full_text.strip().upper()
          if not re.search(r"\b(?:PAGE|SUB)\s*TOTAL\b", clean_upper):
            statement_ended = True
            break
          continue

        if is_closing_balance_line(line_full_text):
          if len(decimal_items) > 0:
            cl_cand = decimal_items[0][0]
            if cl_cand > 0 and in_table_cl_bal is None:
              in_table_cl_bal = cl_cand
              self.log(
                  f"Captured Closing Balance from Bal c/d: ₹{cl_cand:,.2f}"
              )
          continue

        if is_opening_balance_line(line_full_text):
          if len(decimal_items) > 0:
            op_cand = decimal_items[-1][0]
            if op_cand > 0 and in_table_op_bal is None:
              in_table_op_bal = op_cand
              self.log(f"Captured Opening Balance: ₹{op_cand:,.2f}")
          continue

        if any(
            re.search(pat, line_full_text, re.IGNORECASE)
            for pat in FOOTER_AND_TOTAL_PATTERNS
        ):
          continue

        dr_val = 0.0
        cr_val = 0.0
        bal_val = 0.0

        if len(decimal_items) >= 3:
          bal_val = decimal_items[-1][0]
          for amt, x_pos in decimal_items[:-1]:
            if abs(x_pos - dr_center) <= abs(x_pos - cr_center):
              dr_val = amt
            else:
              cr_val = amt
        elif len(decimal_items) == 2:
          bal_val = decimal_items[-1][0]
          tx_amt, tx_x = decimal_items[-2]
          if abs(tx_x - dr_center) <= abs(tx_x - cr_center):
            dr_val = tx_amt
          else:
            cr_val = tx_amt
        elif len(decimal_items) == 1:
          amt, x_pos = decimal_items[0]
          d_bal = abs(x_pos - bal_center)
          d_dr = abs(x_pos - dr_center)
          d_cr = abs(x_pos - cr_center)

          if x_pos >= bal_left or (d_bal < d_dr and d_bal < d_cr):
            bal_val = amt
          elif d_dr <= d_cr:
            dr_val = amt
          else:
            cr_val = amt

        particulars_str = " ".join(particulars_tokens).strip()

        if line_date:
          if current_tx:
            raw_txs.append(current_tx)

          current_tx = {
              "date_obj": line_date,
              "raw_date": date_str,
              "ref_no": "",
              "cheque_no": "",
              "particulars_parts": (
                  [particulars_str] if particulars_str else []
              ),
              "dr": dr_val,
              "cr": cr_val,
              "bal": bal_val,
          }
        elif current_tx:
          # Continuation line fusion (prevents phantom rows)
          if current_tx["bal"] == 0.0 and bal_val > 0:
            current_tx["bal"] = bal_val
          if current_tx["dr"] == 0.0 and dr_val > 0:
            current_tx["dr"] = dr_val
          if current_tx["cr"] == 0.0 and cr_val > 0:
            current_tx["cr"] = cr_val

          if particulars_str and len(current_tx["particulars_parts"]) < 6:
            is_header_row = any(
                _match_keyword(particulars_str, HEADER_KEYWORDS[k])
                for k in ["dr", "cr", "bal", "date"]
            )
            if not is_header_row and not is_annexure_or_legend_line(
                particulars_str
            ):
              current_tx["particulars_parts"].append(particulars_str)

    if current_tx:
      raw_txs.append(current_tx)

    return raw_txs, in_table_op_bal, in_table_cl_bal

  def _reconcile(
      self, raw_txs: list[dict], opening_bal: float | None, tol: float = 0.05
  ):
    output = []
    if not raw_txs:
      op = float(opening_bal) if opening_bal is not None else 0.0
      return output, op, op, False

    # Majority-vote account classification
    std_matches = 0
    od_matches = 0

    for i in range(len(raw_txs) - 1):
      b1 = raw_txs[i].get("bal", 0.0)
      b2 = raw_txs[i + 1].get("bal", 0.0)
      if b1 != 0.0 and b2 != 0.0:
        delta = round(b2 - b1, 2)
        dr2 = raw_txs[i + 1].get("dr", 0.0)
        cr2 = raw_txs[i + 1].get("cr", 0.0)

        if abs(delta - (cr2 - dr2)) <= tol:
          std_matches += 1
        if abs(delta - (dr2 - cr2)) <= tol:
          od_matches += 1

    is_od_account = (od_matches > std_matches) and (od_matches >= 3)
    self.log(
        "Account structure:"
        f" {'OVERDRAFT (OD/CC)' if is_od_account else 'STANDARD SAVINGS/CURRENT'}"
        f" (std={std_matches}, od={od_matches})"
    )

    if opening_bal is None:
      first_tx = raw_txs[0]
      f_bal = first_tx["bal"]
      f_dr = first_tx["dr"]
      f_cr = first_tx["cr"]

      if f_bal > 0:
        if is_od_account:
          opening_bal = round(f_bal - f_dr + f_cr, 2)
        else:
          opening_bal = round(f_bal + f_dr - f_cr, 2)
      else:
        opening_bal = 0.0

      self.log(
          f"Auto-calculated Opening Balance from Tx #1: ₹{opening_bal:,.2f}"
      )

    prev_bal = float(opening_bal)
    prev_pdf_bal = float(opening_bal)
    first_date = raw_txs[0]["date_obj"] if raw_txs else None

    output.append({
        "Date": first_date,
        "Ref No": "",
        "Cheque No": "",
        "Particulars": "OPENING BALANCE",
        "Dr": 0.0,
        "Cr": 0.0,
        "Bal": prev_bal,
        "PDF Bal": prev_bal,
        "Status": "VERIFIED",
        "Warning": "",
    })

    running_math_bal = prev_bal

    for tx in raw_txs:
      extracted_dr = tx.get("dr", 0.0)
      extracted_cr = tx.get("cr", 0.0)
      pdf_bal = tx.get("bal", 0.0)

      dr = extracted_dr
      cr = extracted_cr
      warning_msg = ""
      status = "VERIFIED"

      # Delta calculated strictly from source PDF balances
      has_pdf_delta = pdf_bal > 0.0 and prev_pdf_bal > 0.0
      pdf_delta = round(pdf_bal - prev_pdf_bal, 2) if has_pdf_delta else None

      if has_pdf_delta:
        if not is_od_account:
          expected_cr = pdf_delta if pdf_delta > 0 else 0.0
          expected_dr = abs(pdf_delta) if pdf_delta < 0 else 0.0
        else:
          expected_dr = pdf_delta if pdf_delta > 0 else 0.0
          expected_cr = abs(pdf_delta) if pdf_delta < 0 else 0.0

        if dr == 0.0 and cr == 0.0:
          dr = expected_dr
          cr = expected_cr
          warning_msg = "Derived from PDF Bal Δ"
        else:
          extracted_mag = max(dr, cr)
          delta_mag = abs(pdf_delta)

          if abs(extracted_mag - delta_mag) <= tol:
            if (expected_dr > 0 and dr == 0.0 and cr > 0) or (
                expected_cr > 0 and cr == 0.0 and dr > 0
            ):
              dr = expected_dr
              cr = expected_cr
              warning_msg = "Orientation aligned to PDF Bal Δ"
          else:
            # Extracted amount prevails; flag variance against PDF delta
            warning_msg = (
                f"Review: Extracted (Dr {dr:,.2f}, Cr {cr:,.2f}) != PDF Δ"
                f" ({pdf_delta:+,.2f})"
            )
            status = "DIFF"

        tx["dr"] = dr
        tx["cr"] = cr

      if pdf_bal > 0.0:
        row_bal = pdf_bal
        if has_pdf_delta and status != "DIFF":
          expected_movement = (cr - dr) if not is_od_account else (dr - cr)
          if abs(pdf_delta - expected_movement) <= tol:
            status = "VERIFIED"
          else:
            status = "DIFF"
        prev_pdf_bal = pdf_bal
      else:
        if not is_od_account:
          row_bal = round(prev_bal - dr + cr, 2)
        else:
          row_bal = round(prev_bal + dr - cr, 2)

      if is_od_account:
        running_math_bal = round(running_math_bal + dr - cr, 2)
      else:
        running_math_bal = round(running_math_bal - dr + cr, 2)

      output.append({
          "Date": tx["date_obj"],
          "Ref No": tx["ref_no"],
          "Cheque No": tx["cheque_no"],
          "Particulars": " ".join(tx["particulars_parts"]).strip(),
          "Dr": dr,
          "Cr": cr,
          "Bal": row_bal,
          "PDF Bal": pdf_bal,
          "Status": status,
          "Warning": warning_msg,
      })
      prev_bal = row_bal

    return output, float(opening_bal), running_math_bal, is_od_account


def export_to_excel_styled(
    df,
    out_path: str,
    account_name: str = "",
    account_no: str = "",
    is_od_account: bool = False,
):
  from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
  from openpyxl.utils import get_column_letter
  import pandas as pd

  with pd.ExcelWriter(
      out_path,
      engine="openpyxl",
      date_format="DD-MM-YYYY",
      datetime_format="DD-MM-YYYY",
  ) as writer:
    start_row = 1
    if account_name or account_no:
      start_row = 4

    df.to_excel(
        writer, index=False, sheet_name="Ledger", startrow=start_row - 1
    )
    ws = writer.sheets["Ledger"]

    if account_name or account_no:
      ws["A1"] = "Account Holder Name:"
      ws["B1"] = account_name
      ws["A2"] = "Account Number:"
      ws["B2"] = str(account_no)
      label_font = Font(name="Segoe UI", size=10, bold=True, color="1F497D")
      val_font = Font(name="Segoe UI", size=10, bold=True)
      ws["A1"].font = label_font
      ws["A2"].font = label_font
      ws["B1"].font = val_font
      ws["B2"].font = val_font

    header_fill = PatternFill(
        start_color="1F497D", end_color="1F497D", fill_type="solid"
    )
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    total_border = Border(
        top=Side(style="thin", color="000000"),
        bottom=Side(style="double", color="000000"),
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
    )

    header_row_idx = start_row
    first_data_row = header_row_idx + 1
    last_data_row = header_row_idx + len(df)

    for col_idx in range(1, len(df.columns) + 1):
      cell = ws.cell(row=header_row_idx, column=col_idx)
      cell.fill = header_fill
      cell.font = header_font
      cell.alignment = Alignment(horizontal="center", vertical="center")

    for r in range(first_data_row + 1, last_data_row + 1):
      if is_od_account:
        ws.cell(row=r, column=7, value=f"=G{r-1}+E{r}-F{r}")
      else:
        ws.cell(row=r, column=7, value=f"=G{r-1}-E{r}+F{r}")

      ws.cell(
          row=r,
          column=9,
          value=(
              f'=IF(OR(H{r}="", H{r}=0), "VERIFIED",'
              f' IF(ROUND(ABS(G{r}-H{r}), 2)<=0.05, "VERIFIED", "DIFF"))'
          ),
      )

      existing_warning = ""
      df_idx = r - first_data_row
      if "Warning" in df.columns and 0 <= df_idx < len(df):
        val = df.iloc[df_idx]["Warning"]
        if val:
          existing_warning = str(val).replace('"', '""')

      if existing_warning:
        ws.cell(
            row=r,
            column=10,
            value=(
                f'=IF(I{r}="DIFF", "Math diff: " & TEXT(G{r}-H{r}, "#,##0.00"),'
                f' "{existing_warning}")'
            ),
        )
      else:
        ws.cell(
            row=r,
            column=10,
            value=(
                f'=IF(I{r}="DIFF", "Math diff: " & TEXT(G{r}-H{r}, "#,##0.00"),'
                ' "")'
            ),
        )

    for row in ws.iter_rows(
        min_row=first_data_row,
        max_row=last_data_row,
        min_col=1,
        max_col=ws.max_column,
    ):
      for cell in row:
        cell.border = thin_border
        col_name = df.columns[cell.column - 1]
        if col_name == "Date":
          cell.number_format = "DD-MM-YYYY"
          cell.alignment = Alignment(horizontal="center", vertical="center")
        elif col_name in ["Dr", "Cr", "Bal", "PDF Bal"]:
          cell.number_format = "#,##0.00"
          cell.alignment = Alignment(horizontal="right", vertical="center")
        elif col_name in ["Ref No", "Cheque No", "Status"]:
          cell.alignment = Alignment(horizontal="center", vertical="center")
        elif col_name == "Particulars":
          cell.alignment = Alignment(
              horizontal="left", vertical="center", wrap_text=True
          )

    summary_row = last_data_row + 1
    ws.cell(row=summary_row, column=4, value="TOTAL")
    ws.cell(
        row=summary_row,
        column=5,
        value=f"=SUM(E{first_data_row+1}:E{last_data_row})",
    )
    ws.cell(
        row=summary_row,
        column=6,
        value=f"=SUM(F{first_data_row+1}:F{last_data_row})",
    )
    ws.cell(row=summary_row, column=7, value=f"=G{last_data_row}")

    total_font = Font(name="Segoe UI", size=10, bold=True)
    for c in range(1, len(df.columns) + 1):
      cell = ws.cell(row=summary_row, column=c)
      cell.border = total_border
      cell.font = total_font
      if c in [5, 6, 7]:
        cell.number_format = "#,##0.00"
        cell.alignment = Alignment(horizontal="right", vertical="center")
      elif c == 4:
        cell.alignment = Alignment(horizontal="right", vertical="center")

    for col_idx, col_name in enumerate(df.columns, 1):
      col_letter = get_column_letter(col_idx)
      if col_name == "Particulars":
        ws.column_dimensions[col_letter].width = 46
      elif col_name == "Date":
        ws.column_dimensions[col_letter].width = 14
      elif col_name in ["Dr", "Cr", "Bal", "PDF Bal"]:
        ws.column_dimensions[col_letter].width = 16
      elif col_name == "Warning":
        ws.column_dimensions[col_letter].width = 34
      else:
        ws.column_dimensions[col_letter].width = 15

    ws.freeze_panes = ws.cell(row=first_data_row, column=1)


class StatementConverter:

  def __init__(self, log_fn=None):
    self.log = log_fn or (lambda x: None)

  def parse_pdf(
      self,
      pdf_path: str,
      password: str = None,
      first_tx_date: str = "",
      opening_bal: float = None,
      expected_closing_bal: float = None,
  ):
    parser = AnchorStatementParser(log_fn=self.log)
    return parser.parse(
      pdf_path,
      password=password,
      first_tx_date_str=first_tx_date,
      opening_bal=opening_bal,
      expected_closing_bal=expected_closing_bal,
    )